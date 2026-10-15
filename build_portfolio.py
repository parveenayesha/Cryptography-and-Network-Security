from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

ROOT = Path('/mnt/data/CNS_AI_Smart_Course')
OUT = ROOT
NAVY = '1F3864'; BLUE = '2E5496'; TEAL = '1B7F79'; PALE = 'F2F5FB'; AMBER = 'FFF3D4'; GREEN = 'DFF3E8'; WHITE='FFFFFF'; INK='20304A'; MUTED='5C6980'
TITLE = 'AI-Enabled Framework for Outcome-Based Smart Course Design, Automated CO–PO Mapping and Adaptive Content Generation for Cryptography and Network Security'

# --- diagrams ---
fig, ax = plt.subplots(figsize=(11, 3.2))
ax.set_xlim(0, 11); ax.set_ylim(0, 3.2); ax.axis('off')
boxes = [
    (0.15, 1.05, 2.0, 1.25, 'SMART COURSE DESIGN', 'Syllabus → measurable COs\nCalendar → assessment blueprint'),
    (2.45, 1.05, 2.0, 1.25, 'OBE MAPPING', 'CO–PO evidence\nAttainment and quality checks'),
    (4.75, 1.05, 2.0, 1.25, 'ADAPTIVE LEARNING', 'Pre-check → practice tier\nCipher lab and cases'),
    (7.05, 1.05, 1.75, 1.25, 'EVIDENCE', 'Marks, exit tickets\nReflections'),
    (9.1, 1.05, 1.75, 1.25, 'IMPROVE', 'Review gaps\nAdjust next cycle')]
for x,y,w,h,head,body in boxes:
    patch = FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.06,rounding_size=0.10',edgecolor='#'+BLUE,facecolor='#'+PALE,linewidth=1.5)
    ax.add_patch(patch); ax.text(x+w/2,y+h*0.68,head,ha='center',va='center',fontsize=9,fontweight='bold',color='#'+NAVY)
    ax.text(x+w/2,y+h*0.32,body,ha='center',va='center',fontsize=8,color='#'+INK)
for i in range(4):
    x1=boxes[i][0]+boxes[i][2]+0.02; x2=boxes[i+1][0]-0.04
    ax.add_patch(FancyArrowPatch((x1,1.67),(x2,1.67),arrowstyle='-|>',mutation_scale=14,color='#'+TEAL,linewidth=1.5))
ax.text(5.5,2.75,'HUMAN-IN-THE-LOOP COURSE IMPROVEMENT CYCLE',ha='center',fontsize=12,fontweight='bold',color='#'+NAVY)
ax.text(5.5,0.45,'AI drafts and summarises; faculty verifies curriculum alignment, grades, research claims and corrective action.',ha='center',fontsize=9,color='#'+MUTED)
fig.tight_layout(); fig.savefig(ROOT/'docs/framework_diagram.png',dpi=180,bbox_inches='tight'); plt.close(fig)

fig, ax = plt.subplots(figsize=(9, 4.1))
labels = ['CO1','CO2','CO3','CO4','CO5']; cie=[75,68,54,77,60]; see=[72,65,48,70,56]
x = range(len(labels)); width=.34
ax.bar([i-width/2 for i in x], cie, width, label='CIE %')
ax.bar([i+width/2 for i in x], see, width, label='SEE %')
ax.axhline(60, linestyle='--', linewidth=1, label='Illustrative 60% threshold')
ax.set_xticks(list(x),labels); ax.set_ylim(0,100); ax.set_ylabel('Illustrative average (%)'); ax.set_title('Synthetic example only: CIE and SEE by Course Outcome')
ax.legend(ncol=3,fontsize=8); ax.grid(axis='y',alpha=.2); fig.tight_layout(); fig.savefig(ROOT/'docs/simulated_attainment.png',dpi=180,bbox_inches='tight'); plt.close(fig)

# --- doc utilities ---
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def set_keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr(); keep = OxmlElement('w:keepNext'); pPr.append(keep)

def add_page_number(paragraph):
    run = paragraph.add_run(); fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'),'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)

def setup_doc(doc, short_title):
    sec=doc.sections[0]; sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(9.5); normal.font.color.rgb=RGBColor.from_string(INK)
    normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.08
    for sty,size,col in [('Title',28,NAVY),('Heading 1',18,NAVY),('Heading 2',13,BLUE),('Heading 3',10.5,TEAL)]:
        s=styles[sty]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(col)
        s.paragraph_format.space_before=Pt(10 if sty!='Title' else 0); s.paragraph_format.space_after=Pt(5)
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=header.add_run(short_title); r.font.name='Calibri'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=footer.add_run('FDP Capstone Portfolio  |  U23CS704  |  Page '); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string(MUTED); add_page_number(footer)

def add_p(doc,text='',style=None,bold_lead=None,italic=False,align=None):
    p=doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold=True; p.add_run(text[len(bold_lead):])
    else:
        r=p.add_run(text); r.italic=italic
    if align is not None: p.alignment=align
    return p

def add_bullet(doc,text,level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p

def add_num(doc,text):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p

def add_table(doc, headers, rows, widths=None, fontsize=8.2):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'; table.autofit=True
    hdr=table.rows[0]; set_repeat_table_header(hdr)
    for i,h in enumerate(headers):
        c=hdr.cells[i]; c.text=str(h); set_cell_shading(c,NAVY); set_cell_margins(c)
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(0)
            for r in p.runs: r.font.bold=True; r.font.size=Pt(fontsize); r.font.color.rgb=RGBColor.from_string(WHITE)
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            cells[i].text='' if val is None else str(val); set_cell_margins(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            if ri%2: set_cell_shading(cells[i],PALE)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.size=Pt(fontsize)
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return table

def add_bar(doc,title):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(6)
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),NAVY); pPr.append(shd)
    ind=OxmlElement('w:ind'); ind.set(qn('w:left'),'120'); ind.set(qn('w:right'),'120'); pPr.append(ind)
    r=p.add_run(title); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(WHITE); return p

def add_code(doc, text):
    table=doc.add_table(rows=1,cols=1); table.style='Table Grid'; c=table.cell(0,0); set_cell_shading(c,'F5F7FA'); set_cell_margins(c,80,120,80,120)
    c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run(text); r.font.name='Consolas'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(INK)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def add_url_ref(doc, label, url):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(3); p.add_run(label+' '); rr=p.add_run(url); rr.font.color.rgb=RGBColor.from_string(BLUE); rr.underline=True

def section(doc,n,title,desc=None,newpage=True):
    heading=doc.add_paragraph(style='Heading 1')
    heading.paragraph_format.page_break_before=bool(newpage)
    heading.add_run(f'{n}. {title}')
    if desc: add_p(doc,desc)

# --- Main portfolio ---
doc=Document(); setup_doc(doc,'AI-Enabled CNS Smart Course Portfolio')
# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(58)
r=p.add_run('FACULTY DEVELOPMENT PROGRAMME\nCAPSTONE PROJECT PORTFOLIO'); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=RGBColor.from_string(BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(28)
r=p.add_run(TITLE); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string(NAVY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(20)
r=p.add_run('Course redesigned: U23CS704 — Cryptography and Network Security'); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(TEAL)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(24)
for line in ['Submitted by: [Faculty Name]','Designation / Department: [Assistant Professor, Department of CSE]','Institution: [Institution name]','Academic Year: 2026–27 (tentative course plan)','GitHub Website: https://YOUR-USERNAME.github.io/cns-smart-course/','Date of Submission: [Insert date]']:
    rr=p.add_run(line+'\n'); rr.font.size=Pt(10.5); rr.font.color.rgb=RGBColor.from_string(INK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(38)
r=p.add_run('Based on the uploaded CNS syllabus and FDP capstone format. Draft mappings, sample data and sample assessment are clearly labelled for faculty validation.'); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(MUTED)

# Declaration and overview
doc.add_page_break(); doc.add_heading('Declaration and evidence status',level=1)
add_p(doc,'This portfolio uses AI-assisted drafting as a starting point. The course coordinator must verify every outcome, CO–PO correlation, assessment item, research reference, data calculation and policy before academic adoption. The course title, code, prerequisites, credit, CIE/SEE split, objectives, outcomes and unit topics were taken from the supplied U23CS704 syllabus extract.')
add_p(doc,'Illustrative data statement: the attendance, marks, exit-ticket responses and attainment percentages included in the workbook are synthetic examples intended to demonstrate the workflow. They are not records of actual students and must be replaced with approved class data after course delivery.')
add_p(doc,'Question-paper audit status: no original CNS question paper was attached. The sample paper in this portfolio is a proposed teaching assessment, not a substitute for the source-paper audit required by the blooms-quality-and-rebalance-report skill. The formal seven-criteria quality report and marked rebalanced paper remain pending receipt of the actual CNS question paper.')
add_bar(doc,'Portfolio snapshot')
add_table(doc,['Element','Status / output'],[
    ['Course planning','Course information, syllabus map, proposed measurable COs, draft CO–PO mapping, 15-week plan and AI prompts.'],
    ['Website','Responsive static site source included; live GitHub Pages URL must be created by the account owner.'],
    ['Experiential learning','College portal / public Wi-Fi incident case, design brief, discussion prompts and quiz.'],
    ['Interactive resource','Browser-based Caesar-cipher activity and knowledge-check quiz.'],
    ['Assessment quality','Draft 50-mark sample paper and unit/Bloom/CO blueprint; formal source-paper audit pending.'],
    ['Automation / OBE','Google Apps Script starter, workbook templates, formula-based attainment example and improvement plan.'],
    ['AI agent / research / policy','CourseBuddy architecture, five recent papers, future research question and responsible-AI policy.'],
], widths=[1.55,5.6])

doc.add_heading('Executive summary',level=1)
add_p(doc,'The portfolio redesigns U23CS704 Cryptography and Network Security into an AI-enabled, outcome-based course package. It preserves the syllabus scope across five units while introducing measurable draft learning targets, a weekly plan, scenario-based learning, a browser-based cipher activity, an assessment blueprint, lightweight record automation and a continuous improvement loop. A course website is provided as plain HTML, CSS and JavaScript so it can be hosted on GitHub Pages without a build step or third-party library.')
add_p(doc,'The design follows three principles: AI proposes; faculty approves; evidence drives improvement. AI tools can help draft explanations, generate practice scenarios and summarise anonymous feedback, but must not autonomously determine grades or make high-stakes decisions. The CO–PO matrix and attainment thresholds are draft assumptions because the uploaded CNS syllabus does not define an official matrix or institutional attainment rule.')
doc.add_heading('Human-in-the-loop improvement cycle',level=2)
doc.add_picture(str(ROOT/'docs/framework_diagram.png'),width=Inches(6.35))

# Section 1
section(doc,1,'Course Planning','Course facts below distinguish the exact syllabus details from draft improvements proposed for the capstone.')
doc.add_heading('1.1 Course Information Sheet',level=2)
add_table(doc,['Field','Course information'],[
 ['Course code / title','U23CS704 — Cryptography and Network Security'],
 ['Course type','Core Course'],['Programme','B.E. (CSE), as shown in the uploaded syllabus'],
 ['Credits','3'],['Contact hours','L = 3 per week; T and P are blank in the supplied table. The plan uses 3 lecture hours/week.'],
 ['CIE / SEE','40 marks / 60 marks (as printed in the syllabus)'],
 ['Prerequisites','CN, IS (Computer Networks and Information Security)'],
 ['Course objectives','Familiarize students with network security and cryptography; understand security services, attacks and mechanisms for TCP/IP; comprehend/apply network-, transport- and web-security protocols; understand authentication services and algorithms; build understanding of major network-security issues.'],
 ['Teaching approach','Short concept inputs, worked examples, threat-mapping, guided algorithm comparisons, protocol sequence diagrams, case discussion, cipher lab, retrieval quiz and reflection.'],
 ['Planning assumption','15-week calendar × 3 hours/week = 45 lecture hours. The syllabus extract does not prescribe semester dates or unit-wise hours; adjust against the official calendar.'],
],widths=[1.6,5.5])

doc.add_heading('1.2 Supplied syllabus: five-unit map',level=2)
add_table(doc,['Unit','Topics from the supplied syllabus','Planning focus'],[
 ['I','Security concepts: introduction, need, approaches, principles, attacks, services, mechanisms, network-security model. Cryptography: plaintext/ciphertext, substitution/transposition, encryption/decryption, symmetric/asymmetric key cryptography.','Define threat and protection goals; apply a classical cipher; map threats to services and mechanisms.'],
 ['II','Symmetric/public-key ciphers: block cipher principles, DES, AES, Blowfish, RC5, IDEA, block-cipher operation, stream ciphers, RC4; asymmetric cipher principles, RSA, Diffie–Hellman key exchange.','Compare algorithm roles and limitations; solve carefully scoped RSA/Diffie–Hellman examples.'],
 ['III','Cryptographic hash and RSA algorithm/functions as named in the unit heading; message authentication, SHA-512, authentication requirements, HMAC, CMAC, digital signatures; key distribution, public-key distribution, Kerberos, X.509.','Distinguish confidentiality, integrity, authentication and non-repudiation; design a key/authentication sequence.'],
 ['IV','Web-security considerations, SSL/TLS, HTTPS, SSH, wireless network security, mobile-device security, IEEE 802.11 WLAN and IEEE 802.11i WLAN security.','Explain TLS and certificate checks; assess Wi-Fi/mobile risks; select transport/wireless controls.'],
 ['V','E-mail security: PGP, S/MIME. IP security: overview, architecture, Authentication Header, Encapsulating Security Payload, combining security associations.','Choose e-mail protections and design an IPsec/VPN security configuration for a stated scenario.'],
],widths=[.45,3.5,3.1],fontsize=7.7)

doc.add_heading('1.3 Course Outcomes and Bloom levels',level=2)
add_p(doc,'The five official Course Outcomes below are reproduced in meaning from the uploaded syllabus. The operational verb and Bloom levels are proposed to make teaching and assessment more observable; they need approval from the course coordinator and department.')
add_table(doc,['CO','Official outcome intent from the syllabus','Proposed measurable interpretation','Draft Bloom level'],[
 ['CO1','Understand common information and network threat sources.','Identify and classify typical information/network threats and their likely impact.','K2 Understand'],
 ['CO2','Determine appropriate mechanisms for protecting the network.','Select a relevant security service, mechanism or cryptographic algorithm for a described threat and justify the selection.','K3 Apply / K5 Evaluate'],
 ['CO3','Design a security solution for an application system.','Design an end-to-end security solution that combines suitable cryptographic, authentication and protocol controls.','K6 Create'],
 ['CO4','Understand issues and apply related concepts to protection and communication privacy.','Apply cryptographic and network-security concepts to protect communication privacy, integrity and confidentiality.','K3 Apply'],
 ['CO5','Comprehend various network-security threats.','Analyse threats across network, transport, wireless, e-mail and IP-security contexts and compare mitigations.','K4 Analyse'],
],widths=[.5,2.25,3.4,1.0],fontsize=7.8)

doc.add_heading('1.4 Draft CO–PO mapping',level=2)
add_p(doc,'Scale: 3 = substantial, 2 = moderate, 1 = slight, 0 = no proposed mapping. PO references use the commonly used NBA graduate-attribute numbering; the supplied CNS syllabus does not include an official PO/PSO matrix. The workbook contains a complete draft matrix and should be checked against the institution-approved outcome wording.')
rows=[]
mapvals={
'CO1':[('PO1',3,'Security concepts and threat sources use core engineering knowledge.'),('PO2',3,'Learners classify a threat and analyse its likely impact.'),('PO6',2,'Security incidents affect users and society.'),('PO8',2,'Threat reporting and privacy involve ethical responsibility.'),('PO12',1,'New threats require continuing learning.')],
'CO2':[('PO1',3,'Algorithm properties rely on computing and mathematics.'),('PO2',3,'Selecting a control requires threat and constraint analysis.'),('PO3',2,'The chosen control contributes to system design.'),('PO5',3,'Cryptographic tools and protocol options are compared.'),('PO12',1,'Security practice changes over time.')],
'CO3':[('PO1',3,'Solution design applies core network-security knowledge.'),('PO2',3,'Threats and constraints must be analysed.'),('PO3',3,'The CO explicitly requires designing a security solution.'),('PO4',2,'A design can be tested against stated scenarios.'),('PO8',2,'Controls must consider privacy and ethical use.')],
'CO4':[('PO1',3,'Applies crypto and networking concepts.'),('PO2',2,'Students diagnose privacy/security requirements.'),('PO3',2,'Controls contribute to secure communication design.'),('PO5',3,'TLS/SSH/wireless settings involve modern tools.'),('PO8',2,'Communication privacy has ethical implications.')],
'CO5':[('PO1',3,'Threat analysis requires security knowledge.'),('PO2',3,'Students compare threats and mitigations.'),('PO4',2,'Claims may be tested against a scenario or traffic evidence.'),('PO6',2,'Network threats can affect organisations and society.'),('PO8',2,'Security response should be ethical and authorised.')]
}
for coid, vals in mapvals.items():
    for po,lev,reason in vals: rows.append([coid,po,lev,reason])
add_table(doc,['CO','PO','Level','Evidence / rationale'],rows,widths=[.55,.65,.55,5.4],fontsize=7.8)
add_p(doc,'CO–PO closure rule: each cell should be justified by the CO verb, taught content and assessed evidence. AI-generated correlations are proposals only; faculty must validate each cell. Do not treat an unreviewed similarity score as proof of attainment.')

doc.add_heading('1.5 Semester Work Plan',level=2)
add_table(doc,['Weeks','Unit / focus','Hours*','Teaching-learning methods','AI-enabled activity','Assessment evidence'],[
 ['1–3','Unit I: security concepts, classical cipher','9','Guided explanation, threat mapping, think-pair-share','Pre-check; tiered explanations; Caesar cipher lab','Exit ticket; short quiz; threat-control sheet'],
 ['4–6','Unit II: block/stream and public-key ciphers','9','Worked numerical examples, algorithm comparison','Generate additional practice problems; faculty checks solutions','Assignment; RSA/ Diffie–Hellman worksheet'],
 ['7–9','Unit III: hashes, authentication, signatures, key management','9','Sequence diagrams, compare-and-contrast, case discussion','Adaptive HMAC/CMAC/digital-signature prompts','Case submission; quiz; exit ticket'],
 ['10–12','Unit IV: TLS/HTTPS/SSH and wireless security','9','Protocol walkthrough, browser certificate demo, wireless case','Generate scenario variants for Foundation/Core/Advanced tiers','Protocol worksheet; scenario response'],
 ['13–15','Unit V: PGP/S/MIME/IPsec; integration and review','9','Design brief, peer review, oral defence','AI rubric feedback on anonymised design text; faculty confirms','Integrated security design; reflection and mock assessment'],
],widths=[.55,1.3,.45,1.8,1.7,1.6],fontsize=7.3)
add_p(doc,'*Provisional allocation: 3 hours/week for 15 weeks; equal 9-hour unit blocks are a planning convenience, not prescribed syllabus weightage. The official timetable and holidays take precedence.')

doc.add_heading('1.6 Provisional Course Calendar',level=2)
calendar=[
['1','Unit I intro, need/principles','Threat-source warm-up','CO1'],['2','Attacks, services, mechanisms, network-security model','Threat-to-control match','CO1, CO2'],['3','Classical ciphers; symmetric/asymmetric concepts','Caesar cipher lab + exit ticket','CO2'],['4','Block-cipher principles, DES, AES','Algorithm comparison','CO2'],['5','Blowfish, RC5, IDEA, stream ciphers / RC4','Where-to-use discussion','CO2'],['6','RSA and Diffie–Hellman','Worked key-exchange problems','CO2, CO4'],['7','Hashing, SHA-512, message authentication','Integrity/authentication mini case','CO3, CO4'],['8','HMAC, CMAC, digital signatures','Authenticated message-flow design','CO3'],['9','Key distribution, Kerberos, X.509','Authentication sequence drawing','CO3'],['10','Web security, TLS/SSL, HTTPS','Certificate validation demonstration','CO4'],['11','SSH, mobile and wireless security, 802.11i','Public-Wi-Fi risk exercise','CO4, CO5'],['12','Units I–IV consolidation','Formative mock test and feedback','CO1–CO4'],['13','PGP and S/MIME','Select e-mail protection controls','CO3, CO4'],['14','IPsec overview, AH, ESP, security associations','VPN/IPsec design brief','CO3, CO5'],['15','Integration, recap, reflective review','Secure-system pitch and exit survey','All COs']]
add_table(doc,['Week','Topic','Activity / assessment','CO'],calendar,widths=[.5,2.8,3.0,.8],fontsize=7.5)

doc.add_page_break(); doc.add_heading('1.7 AI prompts used for planning',level=2)
for prompt in [
'P1 — “Use only the attached U23CS704 Cryptography and Network Security syllabus. Extract course facts, objectives, the five official COs and all topics in Units I–V. Preserve source terminology and flag missing items instead of guessing.”',
'P2 — “Draft five measurable CNS learning outcomes aligned to the provided official CO intent. Recommend Bloom levels K1–K6, map each to units, and clearly label every verb or level change as proposed pending faculty validation.”',
'P3 — “Create a provisional 15-week CNS teaching plan using 3 lecture hours per week. Cover all five supplied units, include active-learning activities, formative evidence and CO tags. State that unit-wise hours and semester dates are not stated in the syllabus.”',
'P4 — “Propose a CO–PO correlation matrix on the 1–3 scale with a specific evidence-based reason for each non-zero cell. Identify which mappings are assumptions and ask faculty to verify official PO/PSO wording.”',
'P5 — “Generate a responsive static course website for U23CS704 with course overview, outcomes, five syllabus units, weekly schedule, assessment plan and contact placeholders. Use HTML/CSS/JavaScript only, no external libraries, with an accessible Caesar-cipher lab and quiz.”',
'P6 — “Develop a cybersecurity experiential case using a fictional college portal and public Wi-Fi. Ask learners to identify threats, map controls, state assumptions and justify their recommendations. Avoid instructions that enable attacks on real systems.”',
'P7 — “Write a 50-mark sample CNS assessment for practice with equal unit representation and tagged CO/Bloom levels. Clearly state that duration and pattern are assumptions because the syllabus specifies only CIE 40 and SEE 60.”',
'P8 — “Prepare a CO attainment workbook design with formulas for weighted CIE/SEE results, a transparent 0–3 scale, a review flag and an improvement action. Use only simulated, labelled demonstration data until actual class records are available.”',
'P9 — “Summarise the following verified research references in 2–3 sentences each. Do not invent citation metadata; connect each result to the supplied CNS units and identify one future research question.”',
'P10 — “Create a course-support AI agent specification for FAQs, reminders, progress checks and resource sharing. Include escalation to a human, data minimisation, access controls, audit logs and a rule that AI cannot assign final grades.”'
]: add_bullet(doc,prompt)

# Section 2
section(doc,2,'Course Website and GitHub Pages','The source website is delivered as plain static files in the accompanying ZIP package. The GitHub Pages URL remains a placeholder until it is published from the course coordinator’s account.')
doc.add_heading('2.1 Website components',level=2)
add_table(doc,['File / component','Function'],[
 ['index.html','Course overview, draft COs, complete syllabus map, 15-week schedule, assessment plan, responsible-AI policy, references and contact placeholders.'],
 ['styles.css','Responsive navigation, contrast and accessible spacing, outcome cards, unit accordions, assessment cards and tables.'],
 ['script.js','Mobile navigation; Caesar cipher encrypt/decrypt function; three-question knowledge quiz with immediate feedback.'],
 ['README.md','Local run and step-by-step GitHub Pages deployment; notes on scope and tokens.'],
],widths=[1.6,5.5])
add_p(doc,'The website is static by design. It does not collect form responses, authenticate learners or store student results. Google Forms and Sheets are described as separate, optional tools; real links must be added only after the faculty member creates those resources.')
doc.add_heading('2.2 Publish step by step',level=2)
for step in [
'Open https://github.com/ and sign in or create a GitHub account.',
'Create a public repository named cns-smart-course (or a name approved by the department).',
'Upload index.html, styles.css, script.js and README.md into the repository root and commit the files to the main branch.',
'Open Settings → Pages. Under Build and deployment, choose Deploy from a branch, branch main, folder /(root), then save.',
'Wait for the Pages deployment workflow to complete. The expected URL pattern is https://YOUR-USERNAME.github.io/cns-smart-course/. Replace the placeholder username with the actual account name.',
'Open the live site on desktop and mobile. Test the menu, unit accordions, cipher encrypt/decrypt and quiz.',
'Edit the coordinator, institution, e-mail and office-hour placeholders before sharing with students.',
'Capture a screenshot with the real github.io address visible and insert the actual URL in the portfolio. Do not claim the site is live until it has been published and checked.'
]: add_num(doc,step)
doc.add_heading('2.3 Authentication and credentials: website vs GitHub',level=2)
add_table(doc,['Area','How it works in this project','Safety rule'],[
 ['Website visitor','Browser requests the public static HTML/CSS/JS files from GitHub Pages. Cipher and quiz processing remains in the browser.','No login is required for a normal public Pages site; do not collect passwords or student records in this static page.'],
 ['Application authentication','Not implemented. There is no auth endpoint, session middleware, database, cookie, access token or refresh token in the generated site.','Do not describe the site as having a login system. Add a trusted backend only if authentication is a real requirement.'],
 ['Publishing identity','The publisher signs into GitHub and needs write permission to commit. GitHub authorises the repository operation.','Use account security and least-privilege repository permission. Never publish a personal access token.'],
 ['Future automation','A GitHub Actions workflow or API integration may require credentials only if that capability is added.','Store secrets as GitHub Actions secrets; never hard-code keys in front-end files or commit history.'],
],widths=[1.15,3.25,2.75],fontsize=7.7)
add_p(doc,'No repository URL was supplied for inspection. The authentication notes above describe the generated CNS static site and typical GitHub Pages publishing flow; they are not a code audit of an existing external repository.')
add_p(doc,'Official GitHub Pages, account-authentication and Actions-secrets documentation is listed in the references. These resources describe platform security; the generated static website itself has no login or token flow.')

# Section 3
section(doc,3,'Experiential Learning Module','Case scenario: college portal access over public Wi-Fi.')
doc.add_heading('3.1 Teaching case: “The Exam Notice That Could Have Leaked”',level=2)
add_p(doc,'A fictional engineering college releases examination timetables and department notices through a web portal. During a busy registration week, students connect from a café and a public Wi-Fi hotspot. Some students report browser certificate warnings; a few receive unexpected “reset your password” e-mails. The IT team also discovers that an outdated wireless access point is still used in one building. No actual compromise is assumed in this exercise; the class is asked to conduct a controlled threat assessment and recommend protective controls.')
doc.add_heading('Problem statement',level=3)
add_p(doc,'Design a practical security plan that protects portal confidentiality, data integrity, user authentication and communication privacy while remaining usable for students and staff. The plan must address the web session, identity/certificate validation, e-mail notices and campus/public wireless risks. Learners must justify each control by naming the threat it mitigates and the security property it supports.')
doc.add_heading('Learner deliverables',level=3)
for x in ['A threat table with at least four threats, assets, likely impact and security objectives.','A protocol/control map that includes HTTPS/TLS, certificate checks, MFA or strong authentication, wireless settings, and a way to verify official e-mail notices.','A short incident response recommendation explaining what a student should do after a browser certificate warning or suspicious reset e-mail.','A one-page design note with assumptions and trade-offs.']: add_bullet(doc,x)
doc.add_heading('Discussion questions',level=3)
for x in ['Which asset is at risk in each observed event: credentials, content integrity, identity, availability or privacy?','Why is HTTPS alone not proof that a website is legitimate if a user ignores certificate warnings or follows a phishing link?','Compare TLS certificate validation, an e-mail digital signature and HMAC: what does each help establish, and what does it not establish?','Which controls belong on the client device, wireless network, application, and organisation policy?','How can the college monitor abuse while avoiding unnecessary collection of student browsing information?']: add_bullet(doc,x)
doc.add_heading('Facilitator guide / expected ideas',level=3)
add_table(doc,['Risk cue','Conceptual response','CO evidence'],[
 ['Public Wi-Fi / old access point','Use current supported Wi-Fi security, avoid untrusted networks for sensitive tasks, use properly validated HTTPS, patch device/AP firmware.','CO1, CO2, CO4, CO5'],
 ['Certificate warning','Stop and verify the official portal through a trusted route; do not bypass the warning for convenience.','CO2, CO4'],
 ['Password reset e-mail','Verify sender/context, navigate to portal directly, report suspicious mail; use approved password-reset workflow and stronger authentication.','CO1, CO3, CO5'],
 ['Confidential notice','Use authenticated sender/signature mechanisms where appropriate; apply least-privilege access and secure transport.','CO3, CO4'],
],widths=[1.4,4.4,1.1],fontsize=7.6)
doc.add_page_break(); doc.add_heading('Gamified quiz',level=2)
add_table(doc,['Question','Options','Correct answer and explanation'],[
 ['1. What does certificate validation primarily help confirm?','A. Server identity/trust chain; B. Zero vulnerabilities; C. Guaranteed harmless content','A. It helps the client authenticate the server identity; it does not prove that all content is safe.'],
 ['2. Which mechanism can support message integrity and shared-secret authentication?','A. HMAC; B. Plain Base64; C. Compression','A. HMAC authenticates a message using a shared secret and a hash-based construction.'],
 ['3. A student sees a certificate warning. Best next step?','A. Continue; B. Stop and verify through a trusted route; C. Disable validation','B. Do not bypass the warning for sensitive tasks.'],
 ['4. Which is a safe AI practice in the case?','A. Paste student passwords into a chatbot; B. Upload all names and marks; C. Share anonymised themes after policy review','C. Minimise data and remove identifiers before external processing.'],
],widths=[1.8,2.5,2.6],fontsize=7.5)

# Section 4
section(doc,4,'Interactive Learning Resource','The delivered website includes a working Caesar cipher simulation and a short feedback quiz.')
doc.add_heading('4.1 Caesar cipher activity',level=2)
add_p(doc,'Goal: help students model plaintext, ciphertext, a substitution key and reversible encryption/decryption before studying modern block and public-key ciphers. Learners enter a message, select a shift between 0 and 25, choose encryption or decryption, and compare the output. Letter case is preserved; spaces and punctuation stay unchanged.')
add_p(doc,'Safety note: Caesar cipher is intentionally weak and must never be used to protect actual credentials, personal information or institutional communication.')
doc.add_heading('4.2 Core JavaScript logic',level=2)
add_code(doc,"function caesar(text, shift, decrypt) {\n  const amount = ((Number(shift) % 26) + 26) % 26 * (decrypt ? -1 : 1);\n  return String(text).replace(/[a-z]/gi, ch => {\n    const base = ch >= 'a' && ch <= 'z' ? 97 : 65;\n    return String.fromCharCode(\n      ((ch.charCodeAt(0) - base + amount + 26) % 26) + base\n    );\n  });\n}")
doc.add_heading('4.3 Adaptive content tiers',level=2)
add_table(doc,['Tier','Starting evidence','Activity','Next step'],[
 ['Foundation','Pre-check below 50% or learner requests support','Vocabulary cards: plaintext/ciphertext/key; guided Caesar example; labelled walkthrough of confidentiality vs integrity.','Repeat a short quiz and move to Core after demonstrating the basic distinction.'],
 ['Core','Pre-check 50–79%','Encrypt/decrypt unseen strings; compare substitution and transposition; map threats to security mechanisms.','Complete an RSA or authentication sequence problem.'],
 ['Advanced','Pre-check 80%+','Evaluate the limitations of classical ciphers; compare a modern algorithm’s role; design a high-level secure message flow.','Write a short justification and defend trade-offs orally.'],
],widths=[.85,1.5,3.0,1.8],fontsize=7.6)
add_p(doc,'A pre-check score can suggest practice level but must not permanently label a student. Make all tiers available on request, allow reassessment, and have the instructor review any AI-generated explanation before use.')

# Section 5
section(doc,5,'Assessment Design and Quality Assurance','The internal component split and sample assessment are proposals; only the overall CIE/SEE marks are specified in the uploaded syllabus.')
doc.add_heading('5.1 Proposed CIE structure',level=2)
add_table(doc,['Component','Proposed marks','Evidence / safeguards'],[
 ['Mid-test(s)','20','Use balanced unit coverage, CO tags and a faculty-moderated question set.'],
 ['Assignments / applied cases','10','Include RSA/Diffie–Hellman calculations, a message-authentication comparison and an IPsec design brief.'],
 ['Quizzes / retrieval checks','5','Short, low-stakes checks with answer explanations.'],
 ['Reflection / participation evidence','5','Exit tickets, case discussion and an individual explanation; not attendance-only scoring unless policy says so.'],
 ['Total CIE','40','Proposed internal breakdown; confirm institutional assessment regulations.'],
 ['SEE','60','The syllabus confirms 60 marks but does not specify duration or paper structure in the uploaded extract.'],
],widths=[1.55,1.05,4.6])
doc.add_heading('5.2 Sample practice question paper',level=2)
add_p(doc,'U23CS704 — Cryptography and Network Security | Proposed practice assessment | Time: 2 hours (assumed) | Maximum: 50 marks (assumed). Answer all questions. This is newly drafted for capstone illustration and is not the official university examination pattern.')
add_p(doc,'Section A — Short answers (5 × 2 = 10 marks)',bold_lead='Section A — Short answers (5 × 2 = 10 marks)')
for q in [
'Q1. [2 marks — K2 — CO1 — Unit I] Distinguish a security attack from a security service. Give one example of each.',
'Q2. [2 marks — K3 — CO2 — Unit II] In RSA, let p = 11 and q = 13. Compute n and φ(n). (No need to compute a private key.)',
'Q3. [2 marks — K2 — CO3 — Unit III] State two differences between a cryptographic hash and a digital signature.',
'Q4. [2 marks — K2 — CO4 — Unit IV] What does HTTPS add to ordinary HTTP, and why is certificate validation important?',
'Q5. [2 marks — K2 — CO5 — Unit V] State the primary purpose of IPsec Authentication Header (AH) and Encapsulating Security Payload (ESP).'
]: add_bullet(doc,q)
add_p(doc,'Section B — Application and design (5 × 8 = 40 marks)',bold_lead='Section B — Application and design (5 × 8 = 40 marks)')
for q in [
'Q6. [8 marks — K4 — CO1/CO2 — Unit I] A college portal is accessed from public Wi-Fi. Identify four plausible threats, map each to a security objective, and recommend a suitable service or mechanism. Explain each recommendation.',
'Q7. [8 marks — K4 — CO2/CO4 — Unit II] Compare AES and RSA by key type, typical role, performance and key distribution. For a file-sharing system, recommend how they could be used together and justify the design.',
'Q8. [8 marks — K5 — CO3 — Unit III] A department must distribute signed notices and detect message alteration. Evaluate SHA-512, HMAC/CMAC and digital signatures for this case. Propose key management and public-key validation using Kerberos or X.509 concepts where appropriate.',
'Q9. [8 marks — K4 — CO4/CO5 — Unit IV] Explain the main steps in establishing an HTTPS session and describe what SSH is used for. Analyse the risk of ignoring a browser certificate warning and recommend a safer response. Add two campus wireless-LAN protections.',
'Q10. [8 marks — K6 — CO3/CO5 — Unit V] Design e-mail and network-layer security for confidential examination material exchanged between campuses. Use PGP or S/MIME and an IPsec architecture including AH/ESP and security associations. State assumptions and justify the controls.'
]: add_bullet(doc,q)
doc.add_heading('5.3 Initial design quality check for the sample paper',level=2)
add_table(doc,['Unit','Marks','Share','CO focus','Draft cognitive demand'],[
 ['I','10','20%','CO1, CO2','K2 and K4'],['II','10','20%','CO2, CO4','K3 and K4'],['III','10','20%','CO3','K2 and K5'],['IV','10','20%','CO4, CO5','K2 and K4'],['V','10','20%','CO3, CO5','K2 and K6'],['Total','50','100%','All five COs','Remember/understand, apply, analyse, evaluate and create represented']
],widths=[.7,.75,.75,1.4,3.6])
add_p(doc,'What this check shows: every unit receives the same marks in this proposed practice paper; every CO appears at least once; Section B includes scenario analysis and system design in addition to recall. It does not yet establish actual topic coverage % against every syllabus topic or estimate writing-time fairness from an existing paper. Those measures require the intended source question paper and its official duration, total marks, printed unit labels and stated Bloom levels.')
add_bar(doc,'Formal Bloom quality/rebalance audit — awaiting the source question paper')
add_p(doc,'The uploaded files include the CNS syllabus and capstone format, but not an original CNS exam question paper. The supplied blooms-quality-and-rebalance-report skill has a hard input gate: it requires the source question paper before it can score each gradable item on breadth, depth, variety, difficulty, time–mark fairness, thinking-to-inking and actual Bloom demand; compare stated vs actual Bloom levels; compute a composite; and regenerate only weak questions while preserving unit, CO and marks. Therefore no original-to-rebalanced score, weak-item list or marked rebalanced paper is fabricated here. Attach the intended CNS question paper to finish that named report.')

# Section 6
section(doc,6,'Continuous Assessment Automation','The project bundle provides an Apps Script starter and a formula-driven workbook. The online Google Form and Sheet must be created under the instructor’s account; no live links are invented.')
doc.add_heading('6.1 Workbook tabs',level=2)
add_table(doc,['Tab','Purpose','Key fields'],[
 ['Attendance','Attendance register template','Date, learner code, session/unit, present Y/N, evidence note.'],
 ['Marks','Assessment-wise marks','Learner code, assessment, unit, CO, obtained marks, maximum marks, formula-derived percentage.'],
 ['ExitTickets','Learner reflection evidence','Anonymous code, unit, confidence 1–5, what is clear, what remains unclear, question for next class.'],
 ['Reflection_Summary','Weekly faculty reflection','Unit, response count, average confidence, common themes, faculty action.'],
 ['CO_Attainment','CO result summary','CIE%, SEE%, weighted direct %, direct level, indirect score, overall and review flag.'],
 ['CO_PO_Mapping','Draft correlation matrix','Proposed correlation 0–3 and a rationale for each CO/PO row.'],
 ['Assessment_Blueprint','Question mapping','Question, unit, CO, marks, Bloom level and task type.'],
],widths=[1.45,1.7,4.1],fontsize=7.8)
doc.add_heading('6.2 Form and Apps Script workflow',level=2)
for s in [
'Create a Google Form named CNS Weekly Exit Ticket with an anonymous learner code (optional), unit, confidence scale, most clear, still unclear and question for next class.',
'Link the form to a Google Sheet. Use Extensions → Apps Script and paste the provided CNS_Course_Automation.gs.',
'Run setupCnsWorkbook() once and review the requested permissions. The script creates the record tabs and starter formulas.',
'Install the spreadsheet On form submit trigger for onExitTicketSubmit. Submit test responses using demo values and verify that the ExitTickets sheet receives the response.',
'Run summariseExitTickets() to create a non-identifying keyword/theme summary. Instructor reviews the summary and records a teaching action.',
'If using an external AI model for a narrative summary, export aggregated/anonymised themes only after policy review; do not send student names, IDs, raw grades, passwords or sensitive data.'
]: add_num(doc,s)
doc.add_heading('6.3 Sample Apps Script logic',level=2)
add_code(doc,"function onExitTicketSubmit(e) {\n  const ss = SpreadsheetApp.getActiveSpreadsheet();\n  const target = ss.getSheetByName('ExitTickets');\n  const n = e.namedValues || {};\n  const first = keys => {\n    for (const k of keys) if (n[k] && n[k].length) return String(n[k][0]).trim();\n    return '';\n  };\n  target.appendRow([new Date(), first(['Anonymous code']),\n    first(['Unit']), Number(first(['Confidence (1-5)'])) || '',\n    first(['Most clear']), first(['Still unclear']),\n    first(['Question for next class'])]);\n}")

# Section 7
section(doc,7,'Outcome-Based Analysis and Continuous Improvement','The workbook formulas demonstrate how attainment could be reviewed. All percentages in this section are synthetic example values.')
doc.add_heading('7.1 Illustrative CO attainment',level=2)
add_table(doc,['CO','CIE %','SEE %','Direct %*','Direct level','Indirect (0–3)','Overall (0–3)','Status'],[
 ['CO1','75','72','73.2','2','2.6','2.12','Attained'],['CO2','68','65','66.2','2','2.4','2.08','Attained'],['CO3','54','48','50.4','1','1.7','1.14','Review'],['CO4','77','70','72.8','2','2.5','2.10','Attained'],['CO5','60','56','57.6','1','2.0','1.20','Review']
],widths=[.45,.65,.65,.75,.85,1.0,.95,.75],fontsize=7.3)
add_p(doc,'*Direct % = 40% × CIE % + 60% × SEE %. Direct attainment level example: 0 below 50%; 1 from 50% to below 60%; 2 from 60% to below 75%; 3 at or above 75%. Overall example = 80% × direct level + 20% × indirect survey score. Review flag uses an example target of 2.0. These thresholds and weights are placeholders and must be replaced with the institution-approved attainment policy.')
doc.add_picture(str(ROOT/'docs/simulated_attainment.png'),width=Inches(4.7))
doc.add_heading('7.2 Interpretation of the synthetic example',level=2)
add_p(doc,'CO3 is the weakest illustrated outcome: the sample CIE average is 54%, SEE average 48%, and overall level 1.14. CO5 is also flagged for review at 1.20. Because the values are simulated, these findings demonstrate the workflow only; they are not a claim about actual learner performance.')
doc.add_heading('7.3 Improvement plan',level=2)
add_table(doc,['Signal / CO','Potential learning gap','Next-cycle action','Evidence to check'],[
 ['CO1','Confusion between threat, attack, service and mechanism.','Start Unit I with a concept map and classification quiz.','Quiz analysis; exit-ticket confidence.'],
 ['CO2','Difficulty with algorithm choice or key-exchange steps.','Model RSA/Diffie–Hellman, then remove worked steps.','Numerical marks and learner explanation.'],
 ['CO3','Definitions understood, integrated design remains weak.','Use a guided case, peer critique and short oral defence.','Design rubric: coverage, correctness, justification.'],
 ['CO4','Protocol and certificate steps are disconnected.','Use a TLS/certificate sequence diagram and browser case.','Sequence question; scenario reasoning.'],
 ['CO5','Wireless, e-mail and IP controls are studied separately.','Use a multi-layer security design scenario.','Integrated design; moderation feedback.'],
],widths=[.8,1.75,2.8,1.65],fontsize=7.0)

# Section 8
section(doc,8,'AI Agent for Course Management','Proposed name: CNS CourseBuddy. The agent is a design specification only; no public bot endpoint or live integration has been created.')
doc.add_heading('8.1 Supported tasks',level=2)
add_table(doc,['Task','Example student request','Agent response / human boundary'],[
 ['Syllabus FAQ','“What topics are in Unit IV?”','Retrieve the approved course syllabus and quote the relevant topics; cite the source page.'],
 ['Deadline reminder','“When is the case submission due?”','Read instructor-approved calendar data; state the time zone and refer to the official LMS notice.'],
 ['Resource sharing','“Can I revise HMAC?”','Link to approved lecture notes and suggest practice questions; distinguish generated explanations from official material.'],
 ['Assignment tracking','“Did you receive my submission?”','Read only an approved status source; do not expose other students’ details. Escalate uncertain or missing status.'],
 ['Learning support','“Explain RSA in simple steps.”','Offer a tiered explanation and self-check question; warn when numerical assumptions are missing.'],
 ['Escalation','“I think the portal is compromised.”','Show the authorised reporting channel and ask the student not to share passwords or secrets in chat.'],
],widths=[1.2,2.0,3.8],fontsize=7.6)
doc.add_heading('8.2 Request flow and safeguards',level=2)
for s in ['Learner sends a query through the approved course channel.','Agent classifies the request: syllabus FAQ, resource, reminder, progress status, security incident or unsupported/high-risk request.','For course questions, retrieve only faculty-approved syllabus/calendar/resources. For account status, use a protected institutional API with least-privilege access; never infer grades from conversational context.','Generate a draft response with source link and a confidence/coverage check. If evidence is missing or stakes are high, route to the instructor.','Log minimal metadata needed for quality review; do not retain secrets, unnecessary personal data or raw sensitive text.','Instructor reviews errors, updates approved FAQs, and monitors recurring misconceptions.']:
    add_num(doc,s)
add_code(doc,"Student → Channel / UI → Intent classification →\n  [Approved course source?] -- yes --> Retrieve syllabus/calendar/resource\n         | no / high stakes                         |\n         v                                         v\n  Ask clarification / human escalation ← Draft grounded response\n                          ↓\n             Log minimal metadata and faculty-review errors")
doc.add_heading('8.3 Token and credential handling if integrated',level=2)
add_p(doc,'For a future backend, authenticate faculty/admin users through the institution-approved identity provider; authenticate service-to-service requests using server-side credentials; validate authorisation before returning assignment or grade status; store secrets in a managed secret store; rotate/revoke tokens; and log access. Tokens must never be embedded in the public GitHub Pages JavaScript. Use synthetic data until consent, retention, security and access controls are approved.')

# Section 9
section(doc,9,'Research Integration','Five recent research items were selected from publisher or institutional pages and linked to the CNS units. These brief summaries are for teaching integration, not a substitute for reading each full paper.')
research_items=[
('Souvatzidaki, K. & Limniotis, K. (2025). “Post-Quantum Key Exchange in TLS 1.3: Further Analysis on Performance of New Cryptographic Standards.” Cryptography, 9(4), 73. DOI: 10.3390/cryptography9040073.','Evaluates post-quantum key encapsulation mechanisms in TLS 1.3 under network conditions including packet loss. Connects Unit II public-key cryptography with Unit IV secure transport and introduces performance/security trade-offs.','https://doi.org/10.3390/cryptography9040073'),
('Zhou, J., Fu, W., Hu, W., Sun, Z., He, T. & Zhang, Z. (2024). “Challenges and Advances in Analyzing TLS 1.3-Encrypted Traffic: A Comprehensive Survey.” Electronics, 13(20), 4000. DOI: 10.3390/electronics13204000.','Surveys analysis of TLS 1.3 encrypted traffic. Supports a balanced discussion of how encryption protects content/privacy while defenders assess suspicious traffic using appropriate signals. Connects Unit IV.','https://doi.org/10.3390/electronics13204000'),
('Qureshi, I. A. & Asghar, S. (2023). “A Systematic Review of the IEEE-802.11 Standard’s Enhancements and Limitations.” Wireless Personal Communications, 131, 2539–2572. DOI: 10.1007/s11277-023-10553-7.','Reviews IEEE 802.11 improvements and limitations, giving a research basis for the wireless LAN and 802.11i topics in Unit IV. Students can compare security requirements with usability and deployment constraints.','https://doi.org/10.1007/s11277-023-10553-7'),
('Azad, M. A., Abdullah, S., Arshad, J., Lallie, H. & Ahmed, Y. (2024). “Verify and trust: A multidimensional survey of zero-trust security in the age of IoT.” Internet of Things, 27, 101227. DOI: 10.1016/j.iot.2024.101227.','Explains zero-trust patterns such as continuous verification and least-privilege access in IoT. Although not an explicit heading in the supplied syllabus, it extends Unit I security principles and Units III–IV authentication/network access.','https://doi.org/10.1016/j.iot.2024.101227'),
('Glisic, S. (2024). “Quantum vs post-quantum security for future networks: Survey.” Cyber Security and Applications, 2, 100039. DOI: 10.1016/j.csa.2024.100039.','Surveys conventional, post-quantum and quantum-related security directions. Helps contextualise the syllabus’s RSA and Diffie–Hellman topics and motivates long-term cryptographic transition planning.','https://doi.org/10.1016/j.csa.2024.100039')]
for i,(cite,summary,url) in enumerate(research_items,1):
    doc.add_heading(f'9.{i} {cite.split(".")[0]}',level=2); add_p(doc,cite); add_p(doc,summary); p=doc.add_paragraph(); p.add_run('Source: ').bold=True; r=p.add_run(url); r.font.color.rgb=RGBColor.from_string(BLUE); r.underline=True

doc.add_page_break(); doc.add_heading('9.6 Synthesis and future research direction',level=2)
add_p(doc,'Synthesis: the selected literature points to three connected needs. First, cryptographic protocols must adapt as security assumptions change. Second, security improvements must be evaluated for operational cost and compatibility, not only mathematical strength. Third, network defenders must balance access verification and traffic monitoring with the protection of user privacy. Together, these themes support an applied classroom case linking public-key algorithms, TLS, wireless networks and authentication controls.')
add_p(doc,'Future research idea: evaluate privacy-preserving quantum-readiness for a university network. Inventory systems using RSA/Diffie–Hellman, classify long-lived sensitive data, and test a hybrid post-quantum TLS configuration only in a non-production lab. Measure latency, bandwidth, resource use, compatibility and migration effort using synthetic or authorised data.')

# Section 10
section(doc,10,'Responsible AI Use and Academic Integrity','AI use must support learning while keeping faculty responsible for the curriculum, assessment decisions and safety boundaries.')
doc.add_heading('10.1 Student AI usage policy',level=2)
for s in [
'Permitted with disclosure: ask for simplified explanations, generate extra practice questions, brainstorm a case-study plan, check grammar, and seek formative feedback on a draft when the instructor permits it.',
'Required verification: validate cryptographic calculations, protocol statements, code, standards, and references against the prescribed textbook, official documentation or an instructor-approved source.',
'Not permitted: submit generated work as entirely self-authored where disclosure is required; fabricate citations or lab results; use AI to impersonate a peer; use AI to solve a controlled assessment in violation of the rules; or paste passwords, tokens, identifiable student records or confidential institutional data into public tools.',
'Every submission using AI should disclose the tool, purpose, representative prompt, material contribution and verification steps.',
'Cybersecurity boundaries: practice only on approved toy examples, sandboxes or systems with written authorisation. Do not scan, exploit, intercept or access a real system just because an AI tool suggests it.',
'Faculty maintains a human-in-the-loop process and provides an alternate pathway when students cannot or should not use an AI tool.'
]: add_bullet(doc,s)
doc.add_heading('10.2 AI-resilient assessment strategy',level=2)
add_table(doc,['Design choice','Why it helps'],[
 ['Unseen scenario variations','Students must transfer concepts to a new, specific threat model rather than repeat generic text.'],
 ['Show working and assumptions','Cryptographic calculations, protocol flows and threat-to-control reasoning remain inspectable.'],
 ['Short oral defence / viva','Students explain decisions, limitations, and how they verified AI suggestions.'],
 ['Staged submissions / version history','Makes reasoning and revision visible without treating AI detectors as proof.'],
 ['Approved offline or in-class tasks','Provides evidence independent of uncontrolled external tool access.'],
 ['Source-grounded questions','Require citations to the syllabus, textbook or official protocol documentation.'],
],widths=[2.0,5.0])
doc.add_heading('10.3 Privacy and governance',level=2)
add_p(doc,'Use data minimisation, pseudonymisation, purpose limitation and controlled retention. Do not send identifiable grades or student responses to a third-party AI platform without an approved legal/institutional basis. Keep API secrets outside source files. Access to records should be role-based and logged. A human instructor reviews all AI-generated summaries before acting on them.')
add_p(doc,'NEP 2020 alignment: the design encourages active, experiential and multidisciplinary learning, skill development, critical thinking and ethical practice, consistent with the broad direction of the Government of India’s National Education Policy 2020. It does not claim formal institutional approval; course and assessment rules remain subject to the department and university.')

# Annexures
section(doc,'A','AI Prompts Used','Prompt list is a reusable record of prompts drafted for this capstone. Actual tool screenshots should be captured by the author after running the prompts in the selected AI tool.')
add_p(doc,'Prompts P1–P10 appear in Section 1.7. Additional prompts for implementation and review:')
for prompt in [
'P11 — “Check the proposed assessment against the supplied CNS unit/topic list. Return a coverage table and flag every topic not represented. Do not claim an official Bloom quality audit unless an actual source question paper and its metadata are provided.”',
'P12 — “Review this Apps Script for data minimisation, triggers, formula errors, error handling and accidental credential exposure. Keep API calls disabled unless an approved secret-handling design is specified.”',
'P13 — “Summarise feedback only after names/IDs and direct identifiers have been removed. Return recurring themes, uncertainty and faculty follow-up suggestions; do not grade students or invent statistics.”',
'P14 — “Give a concise literature synthesis from these five publisher-verified CNS research papers. Use only the citation details supplied and distinguish what each paper studies from the future research idea proposed by the author.”'
]: add_bullet(doc,prompt)

doc.add_heading('B. Screenshot evidence checklist',level=1)
add_p(doc,'Screenshots must be collected from the actual tool/session. Do not replace them with mock screenshots or claim a website is deployed before deployment.')
add_table(doc,['Evidence item','Screenshot to insert','Status'],[
 ['AI-assisted course planning','Prompt and the resulting CO/unit table in the tool used.','To capture by author'],
 ['Website repository','GitHub repository file list with repository name visible.','Pending repository creation'],
 ['Published site','Browser screenshot with actual github.io URL and schedule visible.','Pending publication'],
 ['Interactive resource','Cipher lab output and quiz result in the browser.','Can capture after opening site'],
 ['Google Form / Sheet','Form questions and linked tabs with no real student data visible.','Pending setup'],
 ['Apps Script','Function list and successful test run, with no tokens/secrets shown.','Pending setup'],
],widths=[1.5,4.3,1.2])

doc.add_heading('C. Links to populate after deployment',level=1)
add_table(doc,['Resource','Link field / status'],[
 ['GitHub repository','[Paste actual repository URL after creation]'],
 ['GitHub Pages website','https://YOUR-USERNAME.github.io/cns-smart-course/ (replace placeholder after publishing)'],
 ['Google Form','[Paste actual Form responder URL]'],
 ['Google Sheet / workbook copy','[Paste actual Sheet URL; share only with approved users]'],
 ['Apps Script project','[Paste actual Apps Script project URL; restrict editor access]'],
 ['Official CNS question paper for audit','[Attach the PDF/DOCX/image or paste the file reference for the formal Bloom quality/rebalance report]'],
],widths=[2.0,5.0])

doc.add_heading('D. Reflection (approximately 400 words)',level=1)
reflection=("This capstone helped me view course design as a cycle that connects intended outcomes, learning activities, assessment evidence and improvement rather than as a collection of lecture notes. Starting with the supplied U23CS704 Cryptography and Network Security syllabus, I mapped the five units to proposed measurable learning outcomes and drafted a week-wise plan. The outcome statements and CO–PO correlations remain subject to departmental approval because the syllabus extract gives the official CO intent but not a formal CO–PO matrix or unit-wise lecture-hour allocation. Making those assumptions visible was an important lesson in responsible academic use of AI.\n\nThe website provides a practical route for students to revisit the syllabus and test a classical cipher in the browser. The Caesar cipher is deliberately simple: it gives learners an immediate result while opening discussion about why modern cryptography needs stronger algorithms, key management and authentication. The public-Wi-Fi case encourages students to connect cryptographic concepts with everyday decisions such as validating a certificate, checking a message source and reporting suspicious activity. These activities make learning more experiential and give the instructor evidence of reasoning, not only recall.\n\nThe workbook and Apps Script show how routine records and weekly reflections can be organised. The sample data are synthetic, so the attainment results are demonstrations rather than claims about a real class. In an actual semester, I would use the institution-approved attainment thresholds, protect student identifiers, inspect the formulas and discuss the results with colleagues before deciding on remediation. A low CO should trigger a targeted teaching response and a follow-up check, not simply a label for students.\n\nThe research review also shows why CNS remains a changing subject. Recent work on TLS traffic, wireless security, zero-trust access and post-quantum transition offers current cases for the core syllabus. At the same time, AI-generated summaries can be incomplete or incorrect, so publisher pages and original papers must remain the evidence base. The formal question-paper quality and rebalance audit is still pending because an original CNS question paper was not provided; it would be misleading to invent an original score or claim that weak questions had been replaced.\n\nMy central takeaway is that AI can reduce drafting and organisation effort, but it does not remove academic responsibility. Faculty must verify technical correctness, outcome alignment, fairness, privacy and references; students must disclose use and show their own reasoning. The strongest course design is therefore human-led, evidence-based and improved after each teaching cycle.")
add_p(doc,reflection)

doc.add_heading('E. References',level=1)
refs=[
'LIET(A), B.E. (CSE). U23CS704 Cryptography and Network Security, AICTE Model Curriculum effective from Academic Year 2026–27. Syllabus PDF supplied by the user.',
'Government of India, Ministry of Education. National Education Policy 2020. https://www.education.gov.in/sites/upload_files/mhrd/files/NEP_Final_English_0.pdf',
'Stallings, W. Cryptography and Network Security: Principles and Practice, 6th ed., Pearson Education. (As listed in the supplied syllabus.)',
'Kahate, A. Cryptography and Network Security, 3rd ed., McGraw Hill. (As listed in the supplied syllabus.)',
'Souvatzidaki, K. & Limniotis, K. (2025). Post-Quantum Key Exchange in TLS 1.3: Further Analysis on Performance of New Cryptographic Standards. Cryptography, 9(4), 73. https://doi.org/10.3390/cryptography9040073',
'Zhou, J. et al. (2024). Challenges and Advances in Analyzing TLS 1.3-Encrypted Traffic: A Comprehensive Survey. Electronics, 13(20), 4000. https://doi.org/10.3390/electronics13204000',
'Qureshi, I. A. & Asghar, S. (2023). A Systematic Review of the IEEE-802.11 Standard’s Enhancements and Limitations. Wireless Personal Communications, 131, 2539–2572. https://doi.org/10.1007/s11277-023-10553-7',
'Azad, M. A. et al. (2024). Verify and trust: A multidimensional survey of zero-trust security in the age of IoT. Internet of Things, 27, 101227. https://doi.org/10.1016/j.iot.2024.101227',
'Glisic, S. (2024). Quantum vs post-quantum security for future networks: Survey. Cyber Security and Applications, 2, 100039. https://doi.org/10.1016/j.csa.2024.100039',
'GitHub Docs. GitHub Pages. https://docs.github.com/en/pages',
'GitHub Docs. About authentication to GitHub. https://docs.github.com/en/authentication/keeping-your-account-and-data-secure/about-authentication-to-github',
'GitHub Docs. Using secrets in GitHub Actions. https://docs.github.com/en/actions/security-for-github-actions/security-guides/using-secrets-in-github-actions',
'Google Apps Script documentation. https://developers.google.com/apps-script'
]
for r in refs:
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0
    run=p.add_run(r); run.font.size=Pt(7.6)

main_doc=OUT/'CNS_AI_Enabled_Smart_Course_Portfolio.docx'
doc.save(main_doc)

# Separate sample assessment docx
qp=Document(); setup_doc(qp,'CNS Sample Practice Assessment')
p=qp.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(25)
r=p.add_run('U23CS704'); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(BLUE)
p=qp.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('CRYPTOGRAPHY AND NETWORK SECURITY'); r.bold=True; r.font.size=Pt(20); r.font.color.rgb=RGBColor.from_string(NAVY)
p=qp.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Proposed Sample Practice Assessment').bold=True
add_table(qp,['Time (assumed)','Maximum marks (assumed)','Assessment status'],[['2 hours','50','Practice paper for capstone illustration only; not an official university examination pattern.']],widths=[1.6,1.6,4.1])
add_p(qp,'Instructions: Answer all questions. Show all steps for numerical questions, state assumptions, and use clear protocol diagrams where useful. This sample has equal unit marks (10 per unit). The uploaded syllabus confirms CIE = 40 and SEE = 60 but does not specify examination duration or question pattern.')
qp.add_heading('Section A — Short Answers (5 × 2 = 10 marks)',level=2)
for q in [
'Q1. [2 marks — K2 — CO1 — Unit I] Distinguish a security attack from a security service. Give one example of each.',
'Q2. [2 marks — K3 — CO2 — Unit II] In RSA, let p = 11 and q = 13. Compute n and φ(n). (No need to compute a private key.)',
'Q3. [2 marks — K2 — CO3 — Unit III] State two differences between a cryptographic hash and a digital signature.',
'Q4. [2 marks — K2 — CO4 — Unit IV] What does HTTPS add to ordinary HTTP, and why is certificate validation important?',
'Q5. [2 marks — K2 — CO5 — Unit V] State the primary purpose of IPsec Authentication Header (AH) and Encapsulating Security Payload (ESP).'
]: add_bullet(qp,q)
qp.add_heading('Section B — Application and Design (5 × 8 = 40 marks)',level=2)
for q in [
'Q6. [8 marks — K4 — CO1/CO2 — Unit I] A college portal is accessed from public Wi-Fi. Identify four plausible threats, map each threat to a security objective, and recommend a suitable security service or mechanism. Explain why each recommendation is relevant.',
'Q7. [8 marks — K4 — CO2/CO4 — Unit II] Compare AES and RSA with respect to key type, typical role, performance and key distribution. For a file-sharing system, recommend how the two could be used together and justify your design.',
'Q8. [8 marks — K5 — CO3 — Unit III] A department must distribute signed notices and detect message alteration. Evaluate SHA-512, HMAC/CMAC and digital signatures for this case. Propose key management and public-key validation using Kerberos or X.509 concepts where appropriate.',
'Q9. [8 marks — K4 — CO4/CO5 — Unit IV] Explain the main steps in establishing a secure HTTPS session and describe what SSH is used for. Analyse the risk of ignoring a browser certificate warning and recommend a safer response. Add two protections for a campus wireless LAN.',
'Q10. [8 marks — K6 — CO3/CO5 — Unit V] Design an e-mail and network-layer security plan for confidential examination material communicated between campuses. Use relevant PGP or S/MIME concepts and IPsec architecture including AH/ESP and security associations. State assumptions and explain the controls chosen.'
]: add_bullet(qp,q)
qp.add_heading('Paper design tags',level=2)
add_table(qp,['Question','Unit','CO','Marks','Draft Bloom'],[
 ['Q1','I','CO1',2,'K2'],['Q2','II','CO2',2,'K3'],['Q3','III','CO3',2,'K2'],['Q4','IV','CO4',2,'K2'],['Q5','V','CO5',2,'K2'],['Q6','I','CO1/CO2',8,'K4'],['Q7','II','CO2/CO4',8,'K4'],['Q8','III','CO3',8,'K5'],['Q9','IV','CO4/CO5',8,'K4'],['Q10','V','CO3/CO5',8,'K6'],['Total','I–V','All COs',50,'Mix']
],widths=[.8,.8,1.35,1.0,1.3],fontsize=8)
add_p(qp,'Audit notice: this is a newly authored sample paper for the capstone. The named blooms-quality-and-rebalance-report skill cannot generate the formal original-question-paper audit until the actual CNS question paper is attached. Once attached, the original should be audited on all seven criteria and only items below the skill’s weak threshold regenerated.')
qp_path=OUT/'CNS_Sample_Assessment_Question_Paper.docx'; qp.save(qp_path)
print(main_doc)
print(qp_path)
